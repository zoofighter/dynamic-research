import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 와이드스크린 설정 (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # 컬러 팔레트 정의 (모던 핀테크 / AI 테크 스타일)
    C_BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    C_BG_LIGHT = RGBColor(248, 250, 252)   # Slate 50
    C_CARD_BG = RGBColor(255, 255, 255)    # Pure White
    C_CARD_BORDER = RGBColor(226, 232, 240)# Slate 200
    C_PRIMARY = RGBColor(37, 99, 235)      # Blue 600
    C_SECONDARY = RGBColor(14, 165, 233)   # Sky 500
    C_ACCENT_GREEN = RGBColor(16, 185, 129)# Emerald 500
    C_ACCENT_PURPLE = RGBColor(139, 92, 246)# Violet 500
    C_ACCENT_ORANGE = RGBColor(249, 115, 22)# Orange 500
    C_TEXT_DARK = RGBColor(30, 41, 59)     # Slate 800
    C_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    C_TEXT_LIGHT = RGBColor(241, 245, 249) # Slate 100

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="DYNAMIC RESEARCH · PROJECT ARCHITECTURE"):
        # 카테고리 태그
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf = tag_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_PRIMARY

        # 메인 타이틀
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TEXT_DARK

    def create_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.2)
        else:
            card.line.fill.background()
        return card

    # =========================================================================
    # SLIDE 1: 표지 (Cover)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, C_BG_DARK)

    # 액센트 바
    accent_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(1.8), Inches(0.08))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = C_SECONDARY
    accent_bar.line.fill.background()

    tag_box = s1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(10), Inches(0.5))
    tf1 = tag_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "PROJECT STRUCTURE & REPOSITORY OVERVIEW"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = C_SECONDARY

    title_box = s1.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(11), Inches(1.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    pt0 = tf_t.paragraphs[0]
    pt0.text = "Dynamic Research 시스템 구성 및 현황"
    pt0.font.size = Pt(40)
    pt0.font.bold = True
    pt0.font.color.rgb = RGBColor(255, 255, 255)

    pt1 = tf_t.add_paragraph()
    pt1.text = "코드베이스 폴더 구조, 핵심 설계 문서 및 실행 파이프라인 분석"
    pt1.font.size = Pt(22)
    pt1.font.color.rgb = RGBColor(203, 213, 225)
    pt1.space_before = Pt(10)

    # 프로젝트 정보 카드
    info_card = create_card(s1, Inches(1.2), Inches(4.8), Inches(10.9), Inches(1.7), RGBColor(30, 41, 59), RGBColor(51, 65, 85))
    info_box = s1.shapes.add_textbox(Inches(1.5), Inches(5.0), Inches(10.3), Inches(1.3))
    tf_info = info_box.text_frame
    tf_info.word_wrap = True
    
    p_info = tf_info.paragraphs[0]
    p_info.text = "📂 프로젝트 루트: /Users/boon/Dropbox/03_code/c_1003_dynamic_research"
    p_info.font.size = Pt(15)
    p_info.font.bold = True
    p_info.font.color.rgb = C_TEXT_LIGHT

    p_info2 = tf_info.add_paragraph()
    p_info2.text = "• LangGraph 기반 2-Agent 분리 파이프라인 (TopicOutlineGraph & DLSResearchGraph)\n• 100% 동적 라이브 검색(DLS) 및 Human-in-the-Loop(HITL) 심층 리서치 엔진 구축 현황"
    p_info2.font.size = Pt(13)
    p_info2.font.color.rgb = RGBColor(148, 163, 184)
    p_info2.space_before = Pt(6)

    # =========================================================================
    # SLIDE 2: 전체 디렉토리 맵 & 4대 구성 영역
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, C_BG_LIGHT)
    add_header(s2, "프로젝트 전체 디렉토리 맵 & 4대 핵심 구성 영역")

    col_w = Inches(2.75)
    col_gap = Inches(0.24)
    start_x = Inches(0.8)
    card_y = Inches(1.6)
    card_h = Inches(5.3)

    sections = [
        {
            "tag": "01. 설정 & 환경",
            "tag_color": C_PRIMARY,
            "title": "Config & Env",
            "path": "./config, requirements.txt, .venv",
            "items": [
                ("requirements.txt", "LangGraph, Pydantic, python-pptx 등 필수 패키지 명세"),
                (".venv / .gitignore", "가상환경 격리 및 Git 버전 관리 무결성"),
                ("config/", "모델 프로바이더, API 키, 실행 파라미터 구성"),
                ("README.md", "프로젝트 설치 및 실행 가이드라인 제공")
            ]
        },
        {
            "tag": "02. 소스 코드",
            "tag_color": C_SECONDARY,
            "title": "Source Code (src/)",
            "path": "./src (graphs, nodes, prompts...)",
            "items": [
                ("graphs/", "LangGraph StateGraph 워크플로우 정의"),
                ("nodes/", "주제추출, 검색, 크롤링, 집필 등 노드 로직"),
                ("prompts/", "클러스터링, 다각도 쿼리, 보고서 작성 프롬프트"),
                ("providers/ & utils/", "Serper/DuckDuckGo, Jina Reader, 공통 유틸")
            ]
        },
        {
            "tag": "03. 문서화 체계",
            "tag_color": C_ACCENT_PURPLE,
            "title": "Documentation (docs/)",
            "path": "./docs (총 15+ 핵심 문서)",
            "items": [
                ("설계/기획 문서", "dls_design, project_overview, requirements"),
                ("아키텍처 연구", "langgraph_design, dls_vs_agentic_rag"),
                ("구현 & 검증 계획", "implementation_plan, test_plan, pre_review"),
                ("발표 & 레퍼런스", "presentation_content, PPTX 슬라이드")
            ]
        },
        {
            "tag": "04. 실행 & 파이프라인",
            "tag_color": C_ACCENT_ORANGE,
            "title": "Execution & Output",
            "path": "./scripts, tests, output, temp",
            "items": [
                ("scripts/", "create_presentation.py 등 자동화 스크립트"),
                ("tests/", "노드별 단위 테스트 및 E2E 파이프라인 검증"),
                ("output/", "최종 마크다운 분석 보고서 및 JSON 산출물"),
                ("temp/", "웹 크롤링 캐시 및 중간 파싱 데이터 저장")
            ]
        }
    ]

    for i, sec in enumerate(sections):
        x = start_x + i * (col_w + col_gap)
        create_card(s2, x, card_y, col_w, card_h)

        # 상단 태그
        tag_bg = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), card_y + Inches(0.2), col_w - Inches(0.4), Inches(0.35))
        tag_bg.fill.solid()
        tag_bg.fill.fore_color.rgb = sec["tag_color"]
        tag_bg.line.fill.background()
        tf_tag = tag_bg.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = sec["tag"]
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = RGBColor(255, 255, 255)
        p_tag.alignment = PP_ALIGN.CENTER

        # 카드 타이틀 & 경로
        t_box = s2.shapes.add_textbox(x + Inches(0.2), card_y + Inches(0.65), col_w - Inches(0.4), Inches(0.75))
        tf_tb = t_box.text_frame
        tf_tb.word_wrap = True
        tf_tb.margin_left = tf_tb.margin_top = tf_tb.margin_right = tf_tb.margin_bottom = 0
        p_t = tf_tb.paragraphs[0]
        p_t.text = sec["title"]
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TEXT_DARK

        p_sub = tf_tb.add_paragraph()
        p_sub.text = sec["path"]
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = C_TEXT_MUTED
        p_sub.space_before = Pt(3)

        # 아이템 리스트
        item_y = card_y + Inches(1.5)
        for name, desc in sec["items"]:
            ibox = s2.shapes.add_textbox(x + Inches(0.2), item_y, col_w - Inches(0.4), Inches(0.85))
            tf_i = ibox.text_frame
            tf_i.word_wrap = True
            tf_i.margin_left = tf_i.margin_top = tf_i.margin_right = tf_i.margin_bottom = 0
            pi = tf_i.paragraphs[0]
            pi.text = f"• {name}"
            pi.font.size = Pt(12)
            pi.font.bold = True
            pi.font.color.rgb = C_PRIMARY

            pid = tf_i.add_paragraph()
            pid.text = desc
            pid.font.size = Pt(10.5)
            pid.font.color.rgb = C_TEXT_DARK
            pid.space_before = Pt(2)

            item_y += Inches(0.88)

    # =========================================================================
    # SLIDE 3: 소스 코드 구조 심층 분석 (src/)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, C_BG_LIGHT)
    add_header(s3, "소스 코드(src/) 구조: LangGraph 기반 2-Agent 모듈화")

    # 상단 안내 바
    top_bar = create_card(s3, Inches(0.8), Inches(1.5), Inches(11.73), Inches(0.85), RGBColor(238, 242, 255), RGBColor(199, 210, 254))
    tb_box = s3.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.3), Inches(0.75))
    tf_tb2 = tb_box.text_frame
    tf_tb2.word_wrap = True
    p_tbb = tf_tb2.paragraphs[0]
    p_tbb.text = "src/ 모듈은 2개의 독립적인 StateGraph와 이를 보조하는 Providers, Prompts, Utils로 구성됩니다."
    p_tbb.font.size = Pt(13)
    p_tbb.font.bold = True
    p_tbb.font.color.rgb = C_PRIMARY
    p_tbb_sub = tf_tb2.add_paragraph()
    p_tbb_sub.text = "기획 단계(TopicOutline)와 자율 조사 단계(DLSResearch)가 approved_outline.json 규격을 통해 결합도를 낮추고 모듈 독립성을 극대화합니다."
    p_tbb_sub.font.size = Pt(11.5)
    p_tbb_sub.font.color.rgb = C_TEXT_DARK
    p_tbb_sub.space_before = Pt(3)

    # 3개 컬럼 레이아웃
    w3 = Inches(3.75)
    gap3 = Inches(0.24)
    y3 = Inches(2.55)
    h3 = Inches(4.4)

    src_cols = [
        {
            "title": "src/graphs & nodes",
            "badge": "LangGraph 핵심 엔진",
            "badge_color": C_PRIMARY,
            "items": [
                ("TopicOutlineGraph", "뉴스 스캔 → LLM 클러스터링 → 3개 목차 생성 → Human Interrupt(승인)"),
                ("DLSResearchGraph", "균형 쿼리 생성 → Serper 검색 → Jina Reader 크롤링 → Reflection → 집필"),
                ("State & Router", "TypedDict 기반 불변 상태 관리와 조건부 엣지(Conditional Edge) 제어"),
                ("Human-in-the-Loop", "터미널 및 API 상에서 목차를 최종 승인/수정하는 대화형 체크포인트")
            ]
        },
        {
            "title": "src/providers & utils",
            "badge": "외부 서비스 & 유틸",
            "badge_color": C_SECONDARY,
            "items": [
                ("LLM Providers", "OpenAI GPT-4o / Claude 3.5 Sonnet 연동 및 Fallback 라우팅"),
                ("Search Engine", "Google Serper API + DuckDuckGo 하이브리드 검색기"),
                ("Web Crawler", "Jina Reader API 기반 웹페이지 Clean Markdown 즉시 변환"),
                ("Output Helpers", "마크다운 보고서 자동 생성, 인라인 각주 및 출처 테이블 조립기")
            ]
        },
        {
            "title": "src/prompts",
            "badge": "전문 프롬프트 자산",
            "badge_color": C_ACCENT_GREEN,
            "items": [
                ("주제 발굴 프롬프트", "최신 100건 뉴스에서 시의성/파급력 높은 5개 핫 토픽 추출"),
                ("목차 기획 프롬프트", "독자 수준별(심층형/개요형/비즈니스형) 4단 논리 목차 설계"),
                ("쿼리 확장 프롬프트", "확증 편향을 배제하고 찬반/리스크를 아우르는 3~5개 다각도 질의"),
                ("섹션 집필 프롬프트", "수집 원문의 인라인 인용구([1], [2]) 강제 및 팩트 중심 서술")
            ]
        }
    ]

    for i, col in enumerate(src_cols):
        x = Inches(0.8) + i * (w3 + gap3)
        create_card(s3, x, y3, w3, h3)

        # 뱃지
        badge = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y3 + Inches(0.2), Inches(1.8), Inches(0.32))
        badge.fill.solid()
        badge.fill.fore_color.rgb = col["badge_color"]
        badge.line.fill.background()
        tf_b = badge.text_frame
        pb = tf_b.paragraphs[0]
        pb.text = col["badge"]
        pb.font.size = Pt(10)
        pb.font.bold = True
        pb.font.color.rgb = RGBColor(255, 255, 255)
        pb.alignment = PP_ALIGN.CENTER

        # 컬럼 제목
        ct_box = s3.shapes.add_textbox(x + Inches(0.2), y3 + Inches(0.6), w3 - Inches(0.4), Inches(0.45))
        tf_ct = ct_box.text_frame
        tf_ct.margin_left = tf_ct.margin_top = tf_ct.margin_right = tf_ct.margin_bottom = 0
        pct = tf_ct.paragraphs[0]
        pct.text = col["title"]
        pct.font.size = Pt(16)
        pct.font.bold = True
        pct.font.color.rgb = C_TEXT_DARK

        # 리스트
        iy = y3 + Inches(1.15)
        for iname, idesc in col["items"]:
            ibox = s3.shapes.add_textbox(x + Inches(0.2), iy, w3 - Inches(0.4), Inches(0.72))
            tf_item = ibox.text_frame
            tf_item.word_wrap = True
            tf_item.margin_left = tf_item.margin_top = tf_item.margin_right = tf_item.margin_bottom = 0
            pi = tf_item.paragraphs[0]
            pi.text = f"✔ {iname}"
            pi.font.size = Pt(12)
            pi.font.bold = True
            pi.font.color.rgb = C_TEXT_DARK

            pid = tf_item.add_paragraph()
            pid.text = idesc
            pid.font.size = Pt(10.5)
            pid.font.color.rgb = C_TEXT_MUTED
            pid.space_before = Pt(2)

            iy += Inches(0.76)

    # =========================================================================
    # SLIDE 4: 문서화 체계 분석 (docs/)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, C_BG_LIGHT)
    add_header(s4, "문서화 체계(docs/): 체계적인 설계·연구·검증 아티팩트")

    # 4개 영역 2x2 그리드
    gw = Inches(5.7)
    gh = Inches(2.55)
    gx1 = Inches(0.8)
    gx2 = Inches(6.83)
    gy1 = Inches(1.55)
    gy2 = Inches(4.35)

    doc_groups = [
        (gx1, gy1, "1. 기획 및 요구사항 명세", C_PRIMARY, [
            ("requirements_20261003.md", "DLS 시스템의 기능적/비기능적 요구사항 및 성공 기준"),
            ("project_overview_20261003.md", "프로젝트 추진 배경, 목표 아키텍처 및 마일스톤 로드맵"),
            ("topic_outline_agent_design_20261004.md", "뉴스 스캐닝 및 휴먼 인터럽트(HITL) 아웃라인 상세 설계")
        ]),
        (gx2, gy1, "2. 아키텍처 및 비교 연구", C_ACCENT_PURPLE, [
            ("dls_design_20261003.md", "동적 라이브 검색(DLS) 핵심 엔진 및 다각도 쿼리 아키텍처"),
            ("dls_vs_agentic_rag_20261004.md", "고정 벡터 DB(RAG) vs 동적 실시간 검색(DLS) 심층 비교 분석"),
            ("llamaindex_integration_and_benchmarking.md", "LlamaIndex 연동 방안 및 리서치 성능 벤치마킹 체계")
        ]),
        (gx1, gy2, "3. 구현 및 사전 검토", C_SECONDARY, [
            ("langgraph_design_20261004.md", "LangGraph 상태 전이, 2-Graph 독립 실행 및 체크포인터 구조"),
            ("pre_implementation_review_20261004.md", "구현 전 리스크 점검, Fallback 전략 및 토큰 비용 최적화"),
            ("implementation_plan_20261004.md", "단계별 코드 구현 순서 및 모듈별 작업 체크리스트")
        ]),
        (gx2, gy2, "4. 테스트 및 커뮤니케이션", C_ACCENT_ORANGE, [
            ("test_plan_20261004.md", "단위 테스트(pytest), E2E 통합 테스트 및 환각 검증 지표"),
            ("presentation_content_20261004.md", "자율 리서치 시스템 발표용 시나리오 및 슬라이드 원고"),
            ("dynamic_research_presentation_20261004.pptx", "이해관계자 보고용 16:9 와이드스크린 공식 프레젠테이션 덱")
        ])
    ]

    for gx, gy, gtitle, gcolor, gitems in doc_groups:
        create_card(s4, gx, gy, gw, gh)

        # 헤더 텍스트
        h_box = s4.shapes.add_textbox(gx + Inches(0.25), gy + Inches(0.2), gw - Inches(0.5), Inches(0.4))
        tf_h = h_box.text_frame
        tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
        ph = tf_h.paragraphs[0]
        ph.text = gtitle
        ph.font.size = Pt(15)
        ph.font.bold = True
        ph.font.color.rgb = gcolor

        # 내부 항목
        iy = gy + Inches(0.65)
        for doc_name, doc_desc in gitems:
            dbox = s4.shapes.add_textbox(gx + Inches(0.25), iy, gw - Inches(0.5), Inches(0.55))
            tf_d = dbox.text_frame
            tf_d.word_wrap = True
            tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
            pd = tf_d.paragraphs[0]
            pd.text = f"📄 {doc_name}"
            pd.font.size = Pt(11.5)
            pd.font.bold = True
            pd.font.color.rgb = C_TEXT_DARK

            pdd = tf_d.add_paragraph()
            pdd.text = doc_desc
            pdd.font.size = Pt(10)
            pdd.font.color.rgb = C_TEXT_MUTED
            pdd.space_before = Pt(1)

            iy += Inches(0.58)

    # =========================================================================
    # SLIDE 5: 실행 파이프라인 및 개발 환경 (Scripts, Tests, Config)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, C_BG_LIGHT)
    add_header(s5, "환경 설정, 자동화 스크립트 및 테스트·산출물 파이프라인")

    # 3개 메인 카드
    w5 = Inches(3.75)
    gap5 = Inches(0.24)
    y5 = Inches(1.6)
    h5 = Inches(5.3)

    col_configs = [
        {
            "title": "설정 & 가상환경 (Config & Env)",
            "subtitle": "안정적인 개발 런타임 보장",
            "tag": "Environment",
            "tag_color": C_PRIMARY,
            "blocks": [
                ("Python 3.11+ & .venv", "격리된 가상환경에 최신 의존성 패키지 설치 완료"),
                ("requirements.txt", "langgraph, langchain-openai, python-pptx, pytest 등 관리"),
                ("config/settings.yaml", "검색 API 키, LLM Temperature, 크롤링 타임아웃 중앙화"),
                (".gitignore & README.md", "임시 파일 및 캐시 제외, 협업 표준 온보딩 문서 제공")
            ]
        },
        {
            "title": "자동화 스크립트 & 산출물",
            "subtitle": "스크립트 실행 및 결과 저장 체계",
            "tag": "Scripts & Outputs",
            "tag_color": C_SECONDARY,
            "blocks": [
                ("scripts/create_presentation.py", "python-pptx 기반 발표 슬라이드 코드형 자동 생성(IaC)"),
                ("output/ 디렉토리", "최종 생성된 심층 리서치 마크다운 보고서(run_id.md) 보관"),
                ("output/approved_outline.json", "Graph 1 승인 목차를 저장하여 Graph 2 입력으로 연계"),
                ("temp/ 디렉토리", "중간 크롤링 텍스트, HTML 변환 캐시 등 임시 버퍼")
            ]
        },
        {
            "title": "테스트 & 품질 보증 (QA)",
            "subtitle": "모듈별 검증 및 CI/CD 연동 준비",
            "tag": "Testing & QA",
            "tag_color": C_ACCENT_GREEN,
            "blocks": [
                ("pytest 기반 단위 테스트", "tests/ 디렉토리 내 각 노드별 독립 실행 테스트 구축"),
                ("Mock Provider 지원", "실제 외부 API 과금 없이 오프라인 파이프라인 동작 검증"),
                ("환각 및 출처 검증", "생성 문장과 수집 원문 URL 매핑률 정량적 평가"),
                ("E2E 파이프라인 무결성", "주제 추출부터 보고서 출력까지 원클릭 검증 체계")
            ]
        }
    ]

    for i, c in enumerate(col_configs):
        x = Inches(0.8) + i * (w5 + gap5)
        create_card(s5, x, y5, w5, h5)

        # 상단 태그
        tag_bg = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y5 + Inches(0.2), Inches(1.5), Inches(0.3))
        tag_bg.fill.solid()
        tag_bg.fill.fore_color.rgb = c["tag_color"]
        tag_bg.line.fill.background()
        tf_tb = tag_bg.text_frame
        pt = tf_tb.paragraphs[0]
        pt.text = c["tag"]
        pt.font.size = Pt(10)
        pt.font.bold = True
        pt.font.color.rgb = RGBColor(255, 255, 255)
        pt.alignment = PP_ALIGN.CENTER

        # 타이틀
        tt_box = s5.shapes.add_textbox(x + Inches(0.2), y5 + Inches(0.6), w5 - Inches(0.4), Inches(0.7))
        tf_tt = tt_box.text_frame
        tf_tt.word_wrap = True
        tf_tt.margin_left = tf_tt.margin_top = tf_tt.margin_right = tf_tt.margin_bottom = 0
        ptt = tf_tt.paragraphs[0]
        ptt.text = c["title"]
        ptt.font.size = Pt(15)
        ptt.font.bold = True
        ptt.font.color.rgb = C_TEXT_DARK

        psub = tf_tt.add_paragraph()
        psub.text = c["subtitle"]
        psub.font.size = Pt(10.5)
        psub.font.color.rgb = C_TEXT_MUTED
        psub.space_before = Pt(2)

        # 블록 리스트
        by = y5 + Inches(1.4)
        for btitle, bdesc in c["blocks"]:
            bbox = s5.shapes.add_textbox(x + Inches(0.2), by, w5 - Inches(0.4), Inches(0.85))
            tf_b = bbox.text_frame
            tf_b.word_wrap = True
            tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
            pb = tf_b.paragraphs[0]
            pb.text = f"▶ {btitle}"
            pb.font.size = Pt(12)
            pb.font.bold = True
            pb.font.color.rgb = C_TEXT_DARK

            pbd = tf_b.add_paragraph()
            pbd.text = bdesc
            pbd.font.size = Pt(10)
            pbd.font.color.rgb = C_TEXT_MUTED
            pbd.space_before = Pt(2)

            by += Inches(0.9)

    # =========================================================================
    # SLIDE 6: 프로젝트 현황 요약 및 향후 추진 로드맵 (Summary & Roadmap)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, C_BG_LIGHT)
    add_header(s6, "프로젝트 구축 현황 총괄 요약 & 향후 추진 로드맵")

    # 상단 요약 카드 (3개 지표)
    kw = Inches(3.75)
    kh = Inches(1.3)
    ky = Inches(1.5)
    
    kpis = [
        ("설계 및 문서화 완료율", "100%", "기획·아키텍처·테스트 계획 15+ 건 완료", C_PRIMARY),
        ("아키텍처 설계 방식", "2-Graph", "HITL(주제선정) + DLS(자율리서치) 분리", C_SECONDARY),
        ("코드베이스 준비도", "Ready", "가상환경, 테스트셋, 모듈 뼈대 구축 완료", C_ACCENT_GREEN)
    ]

    for i, (k_label, k_val, k_desc, k_col) in enumerate(kpis):
        kx = Inches(0.8) + i * (kw + gap5)
        create_card(s6, kx, ky, kw, kh)
        kbox = s6.shapes.add_textbox(kx + Inches(0.25), ky + Inches(0.18), kw - Inches(0.5), kh - Inches(0.36))
        tf_k = kbox.text_frame
        tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0
        pk_l = tf_k.paragraphs[0]
        pk_l.text = k_label
        pk_l.font.size = Pt(11)
        pk_l.font.bold = True
        pk_l.font.color.rgb = C_TEXT_MUTED

        pk_v = tf_k.add_paragraph()
        pk_v.text = k_val
        pk_v.font.size = Pt(24)
        pk_v.font.bold = True
        pk_v.font.color.rgb = k_col
        pk_v.space_before = Pt(2)

        pk_d = tf_k.add_paragraph()
        pk_d.text = k_desc
        pk_d.font.size = Pt(9.5)
        pk_d.font.color.rgb = C_TEXT_DARK
        pk_d.space_before = Pt(2)

    # 하단 4단계 로드맵 프로세스 카드
    ry = Inches(3.05)
    rh = Inches(3.85)
    rw = Inches(2.75)
    rgap = Inches(0.24)

    phases = [
        ("Phase 1", "아키텍처 및 환경 준비", "완료 (Current)", C_PRIMARY, [
            "프로젝트 폴더 구조 확립",
            "가상환경 및 라이브러리 세팅",
            "15+ 핵심 기술 설계 문서화",
            "DLS vs Agentic RAG 비교 연구"
        ]),
        ("Phase 2", "Graph 1 모듈 구현", "다음 작업 (Next)", C_SECONDARY, [
            "한경 100건 뉴스 스캔 수집기",
            "LLM 주제 클러스터링 노드",
            "3개 아웃라인 자동 제안기",
            "Human-in-the-Loop 승인 CLI"
        ]),
        ("Phase 3", "Graph 2 리서치 구현", "예정 (Planned)", C_ACCENT_PURPLE, [
            "균형 쿼리 생성 엔진",
            "Serper + Jina Reader 수집기",
            "정보 충분성 Reflection 루프",
            "인라인 각주 마크다운 집필기"
        ]),
        ("Phase 4", "통합 검증 및 고도화", "예정 (Planned)", C_ACCENT_ORANGE, [
            "E2E 통합 파이프라인 실행",
            "출처 정확도 및 환각 평가",
            "LlamaIndex 벤치마킹 연동",
            "비즈니스 보고서 자동 배포"
        ])
    ]

    for i, (ph_id, ph_title, ph_status, ph_color, ph_tasks) in enumerate(phases):
        px = Inches(0.8) + i * (rw + rgap)
        create_card(s6, px, ry, rw, rh)

        # 상단 뱃지
        p_badge = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px + Inches(0.2), ry + Inches(0.2), rw - Inches(0.4), Inches(0.32))
        p_badge.fill.solid()
        p_badge.fill.fore_color.rgb = ph_color
        p_badge.line.fill.background()
        tf_pb = p_badge.text_frame
        ppb = tf_pb.paragraphs[0]
        ppb.text = f"{ph_id} · {ph_status}"
        ppb.font.size = Pt(10.5)
        ppb.font.bold = True
        ppb.font.color.rgb = RGBColor(255, 255, 255)
        ppb.alignment = PP_ALIGN.CENTER

        # 단계 제목
        pt_box = s6.shapes.add_textbox(px + Inches(0.2), ry + Inches(0.65), rw - Inches(0.4), Inches(0.55))
        tf_pt = pt_box.text_frame
        tf_pt.word_wrap = True
        tf_pt.margin_left = tf_pt.margin_top = tf_pt.margin_right = tf_pt.margin_bottom = 0
        ppt = tf_pt.paragraphs[0]
        ppt.text = ph_title
        ppt.font.size = Pt(14)
        ppt.font.bold = True
        ppt.font.color.rgb = C_TEXT_DARK

        # 태스크 리스트
        ty = ry + Inches(1.3)
        for task in ph_tasks:
            tbox = s6.shapes.add_textbox(px + Inches(0.2), ty, rw - Inches(0.4), Inches(0.55))
            tf_t = tbox.text_frame
            tf_t.word_wrap = True
            tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
            ptk = tf_t.paragraphs[0]
            ptk.text = f"✔ {task}"
            ptk.font.size = Pt(11)
            ptk.font.color.rgb = C_TEXT_DARK
            ty += Inches(0.55)

    # 출력 파일 저장
    output_path = "/Users/boon/Dropbox/03_code/c_1003_dynamic_research/docs/project_structure_presentation_20261004.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
